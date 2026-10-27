"""Extract citable passages without a language model.

PDF text comes from PyMuPDF. Office files use the same libraries as the
upstream text-only parser (python-docx, python-pptx, openpyxl). Markdown and
plain text are read directly. Page numbers and headings are kept so retrieval
can cite a file, page, and section.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re

from deeptutor.agent_study.errors import StudyError

TEXT_EXTENSIONS = frozenset({".txt", ".md", ".markdown", ".rst", ".csv", ".log", ".text"})
MARKDOWN_EXTENSIONS = frozenset({".md", ".markdown"})
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | frozenset({".pdf", ".docx", ".pptx", ".xlsx"})

_CHUNK_LIMIT = 1200
_HARD_OVERLAP = 80
_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_MAX_FILE_BYTES = 50 * 1024 * 1024


@dataclass(frozen=True)
class Passage:
    text: str
    page: int | None
    section: str | None


def passages_for(path: Path, *, max_bytes: int = _MAX_FILE_BYTES) -> list[Passage]:
    """Return citable passages for one supported file."""

    if not path.is_file():
        raise StudyError(f"Not a file: {path}", code="not_found")
    size = path.stat().st_size
    if size > max_bytes:
        raise StudyError(
            f"{path.name} is {size} bytes, over the {max_bytes} byte ingest limit.",
            code="too_large",
        )
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _pdf_passages(path)
    if suffix == ".docx":
        return _docx_passages(path)
    if suffix == ".pptx":
        return _pptx_passages(path)
    if suffix == ".xlsx":
        return _xlsx_passages(path)
    if suffix in TEXT_EXTENSIONS:
        return _text_passages(path, markdown=suffix in MARKDOWN_EXTENSIONS)
    raise StudyError(
        f"Unsupported file type {suffix or '(none)'}. "
        f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}.",
        code="unsupported",
    )


def chunk_text(text: str, *, limit: int = _CHUNK_LIMIT) -> list[str]:
    cleaned = text.replace("\x00", "").strip()
    if not cleaned:
        return []
    if len(cleaned) <= limit:
        return [cleaned]
    paragraphs = [part.strip() for part in re.split(r"\n\s*\n", cleaned) if part.strip()]
    if not paragraphs:
        paragraphs = [cleaned]
    chunks: list[str] = []
    buffer = ""
    for paragraph in paragraphs:
        if len(paragraph) > limit:
            if buffer:
                chunks.append(buffer)
                buffer = ""
            chunks.extend(_hard_split(paragraph, limit))
            continue
        candidate = f"{buffer}\n\n{paragraph}" if buffer else paragraph
        if len(candidate) <= limit:
            buffer = candidate
            continue
        if buffer:
            chunks.append(buffer)
        buffer = paragraph
    if buffer:
        chunks.append(buffer)
    return chunks


def _hard_split(text: str, limit: int) -> list[str]:
    pieces: list[str] = []
    start = 0
    while start < len(text):
        end = min(len(text), start + limit)
        pieces.append(text[start:end].strip())
        if end >= len(text):
            break
        start = max(0, end - _HARD_OVERLAP)
    return [piece for piece in pieces if piece]


def _pack(text: str, *, page: int | None, section: str | None) -> list[Passage]:
    body = text.strip()
    if section and body and not body.startswith(section):
        body = f"{section}\n{body}"
    return [
        Passage(text=piece, page=page, section=section)
        for piece in chunk_text(body)
        if piece.strip()
    ]


def _pdf_passages(path: Path) -> list[Passage]:
    try:
        import fitz
    except ImportError as exc:
        raise StudyError(
            "PyMuPDF is required to ingest PDFs (it is a DeepTutor dependency).",
            code="parser_missing",
        ) from exc
    try:
        document = fitz.open(path)
    except Exception as exc:
        raise StudyError(f"Could not open {path.name}: {exc}", code="parse") from exc
    passages: list[Passage] = []
    try:
        if document.is_encrypted and not document.authenticate(""):
            raise StudyError(f"{path.name} is encrypted.", code="encrypted")
        for index, page in enumerate(document, start=1):
            raw = page.get_text("text") or ""
            passages.extend(_pack(raw, page=_page_number(index), section=_pdf_section(raw)))
    finally:
        document.close()
    if not passages:
        raise StudyError(f"{path.name} has no extractable text.", code="empty")
    return passages


def _page_number(index: int) -> int:
    return index


def _pdf_section(text: str) -> str | None:
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if len(lines) >= 2 and len(lines[0]) <= 80:
        return lines[0]
    return None


def _text_passages(path: Path, *, markdown: bool) -> list[Passage]:
    try:
        raw = path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        raw = path.read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise StudyError(f"Could not read {path.name}: {exc}", code="parse") from exc
    if markdown:
        passages = _markdown_passages(raw)
    else:
        passages = _pack(raw, page=None, section=None)
    if not passages:
        raise StudyError(f"{path.name} has no extractable text.", code="empty")
    return passages


def _markdown_passages(raw: str) -> list[Passage]:
    sections: list[tuple[str | None, list[str]]] = [(None, [])]
    for line in raw.splitlines():
        match = _HEADING.match(line)
        if match:
            sections.append((match.group(2).strip(), []))
            continue
        sections[-1][1].append(line)
    passages: list[Passage] = []
    for title, lines in sections:
        passages.extend(_pack("\n".join(lines), page=None, section=title))
    return passages


def _docx_passages(path: Path) -> list[Passage]:
    try:
        from docx import Document
    except ImportError as exc:
        raise StudyError(
            "python-docx is required to ingest .docx files.",
            code="parser_missing",
        ) from exc
    try:
        document = Document(str(path))
    except Exception as exc:
        raise StudyError(f"Could not open {path.name}: {exc}", code="parse") from exc
    passages: list[Passage] = []
    section: str | None = None
    buffer: list[str] = []

    def flush() -> None:
        nonlocal buffer
        passages.extend(_pack("\n".join(buffer), page=None, section=section))
        buffer = []

    for paragraph in document.paragraphs:
        style = ""
        if paragraph.style is not None and paragraph.style.name:
            style = paragraph.style.name
        text = paragraph.text.strip()
        if style.startswith("Heading"):
            flush()
            section = text or section
            continue
        if text:
            buffer.append(text)
    flush()
    if not passages:
        raise StudyError(f"{path.name} has no extractable text.", code="empty")
    return passages


def _pptx_passages(path: Path) -> list[Passage]:
    try:
        from pptx import Presentation
    except ImportError as exc:
        raise StudyError(
            "python-pptx is required to ingest .pptx files.",
            code="parser_missing",
        ) from exc
    try:
        presentation = Presentation(str(path))
    except Exception as exc:
        raise StudyError(f"Could not open {path.name}: {exc}", code="parse") from exc
    passages: list[Passage] = []
    for index, slide in enumerate(presentation.slides, start=1):
        title = _slide_title(slide)
        lines: list[str] = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            for paragraph in shape.text_frame.paragraphs:
                text = paragraph.text.strip()
                if text and text != title:
                    lines.append(text)
        passages.extend(_pack("\n".join(lines), page=index, section=title))
    if not passages:
        raise StudyError(f"{path.name} has no extractable text.", code="empty")
    return passages


def _slide_title(slide) -> str | None:
    title_shape = getattr(slide.shapes, "title", None)
    if title_shape is None:
        return None
    text = getattr(title_shape, "text", "") or ""
    cleaned = text.strip()
    return cleaned or None


def _xlsx_passages(path: Path) -> list[Passage]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise StudyError(
            "openpyxl is required to ingest .xlsx files.",
            code="parser_missing",
        ) from exc
    try:
        workbook = load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        raise StudyError(f"Could not open {path.name}: {exc}", code="parse") from exc
    passages: list[Passage] = []
    try:
        for sheet in workbook.worksheets:
            rows: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [
                    str(cell).strip() for cell in row if cell is not None and str(cell).strip()
                ]
                if cells:
                    rows.append(" | ".join(cells))
            passages.extend(_pack("\n".join(rows), page=None, section=sheet.title))
    finally:
        workbook.close()
    if not passages:
        raise StudyError(f"{path.name} has no extractable text.", code="empty")
    return passages
